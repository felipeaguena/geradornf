import os
import sqlite3
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def format_pct(val):
    if val is None or val == "" or str(val).strip() == "-":
        return "-"
    val_str = str(val).strip()
    if val_str == "NT":
        return "NT (0%)"
    try:
        num = float(val_str)
        return f"{num:.2f}%"
    except Exception:
        return val_str

def exportar_ncm_excel():
    base_dir = os.path.dirname(__file__)
    db_path = os.path.join(base_dir, 'database.db')
    output_excel = os.path.join(base_dir, 'Tabela_NCM_Completa_Backup.xlsx')

    print("Conectando ao banco de dados...")
    conn = sqlite3.connect(db_path)
    c = conn.cursor()

    query = """
        SELECT 
            codigo,
            utrib,
            descricao_utrib,
            descricao,
            aliquota_ii,
            aliquota_ipi,
            aliquota_pis,
            aliquota_cofins,
            regime_pis_cofins,
            data_inicio,
            data_fim,
            (tipo_ato || ' ' || numero_ato || '/' || ano_ato)
        FROM ncm
        ORDER BY codigo ASC
    """

    print("Extraindo todos os registros da tabela NCM...")
    c.execute(query)
    rows = c.fetchall()
    conn.close()

    print(f"Total de registros a exportar: {len(rows)}")
    print(f"Gerando arquivo Excel formatado em: {output_excel}...")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tabela NCM e Impostos"

    headers = [
        "Código NCM",
        "uTrib",
        "Descrição da uTrib",
        "Descrição Oficial do NCM",
        "Alíquota II (%)",
        "Alíquota IPI (%)",
        "Alíquota PIS (%)",
        "Alíquota COFINS (%)",
        "Regime PIS/COFINS",
        "Início de Vigência",
        "Fim de Vigência",
        "Ato Legal"
    ]

    header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    # Inserir Cabeçalho
    ws.append(headers)
    ws.row_dimensions[1].height = 28

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    # Inserir Linhas
    for row_idx, r in enumerate(rows, start=2):
        ii_val = format_pct(r[4])
        ipi_val = format_pct(r[5])
        pis_val = format_pct(r[6])
        cofins_val = format_pct(r[7])

        row_data = [
            r[0] or "",
            r[1] or "",
            r[2] or "",
            r[3] or "",
            ii_val,
            ipi_val,
            pis_val,
            cofins_val,
            r[8] or "",
            r[9] or "",
            r[10] or "",
            r[11] or ""
        ]
        ws.append(row_data)

        # Alinhamentos específicos
        ws.cell(row=row_idx, column=1).alignment = align_left # Código NCM
        ws.cell(row=row_idx, column=2).alignment = align_center # uTrib
        ws.cell(row=row_idx, column=5).alignment = align_center # II
        ws.cell(row=row_idx, column=6).alignment = align_center # IPI
        ws.cell(row=row_idx, column=7).alignment = align_center # PIS
        ws.cell(row=row_idx, column=8).alignment = align_center # COFINS
        ws.cell(row=row_idx, column=10).alignment = align_center # Data Início
        ws.cell(row=row_idx, column=11).alignment = align_center # Data Fim

    # Congelar painel na primeira linha
    ws.freeze_panes = "A2"

    # Larguras otimizadas
    col_widths = {
        1: 15,  # Código NCM
        2: 10,  # uTrib
        3: 20,  # Descrição uTrib
        4: 60,  # Descrição Oficial
        5: 15,  # II
        6: 15,  # IPI
        7: 15,  # PIS
        8: 16,  # COFINS
        9: 32,  # Regime PIS/COFINS
        10: 16, # Início
        11: 16, # Fim
        12: 24  # Ato Legal
    }
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    wb.save(output_excel)
    file_size_mb = os.path.getsize(output_excel) / (1024 * 1024)
    print(f"Sucesso! Novo backup Excel gerado: {output_excel} ({file_size_mb:.2f} MB, {len(rows)} linhas)")
    return output_excel

if __name__ == '__main__':
    exportar_ncm_excel()
